"""Excel (.xlsx) writer for the FCU schedule download.

Produces the template structure that CSV cannot express: row 1 is exactly
two merged cells — "FCU SCHEDULE" spanning A1:F1 and the user's company
logo spanning G1:N1 (no column lines in between) — followed by the four
label/value detail rows (values merged B:F and H:N), a styled column-header
row, and the data.

Every data value is written as the exact extracted string (zero rounding),
same guarantee as the CSV path.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path  # noqa: F401  (used by the logo fallback)

from PIL import Image as PILImage
from PIL import ImageChops
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

from .synthesizer import LOGO_KEY, PROJECT_FIELDS

NAVY = "141B4D"
LOGO_INK = "1F3B57"
GRID = "9A9A9A"

_NUM_COLS = 14
_COL_WIDTHS = [30, 12, 12, 12, 11, 14, 14, 15, 9, 7, 10, 8, 12, 16]

_LOGO_COL = 7           # G — first column of the logo zone (1-based)
_TITLE_ROW_PT = 72      # row 1 height in points (tall header, as in the template)
_PX_PER_WIDTH = 7       # Excel column-width unit -> pixels
_PT_TO_PX = 4 / 3
_LOGO_PAD_PX = 6        # breathing room inside the merged zone
_LOGO_MAX_UPSCALE = 3.0  # enlarge small logos, but not into a blur

_thin = Side(style="thin", color=GRID)
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
_left = Alignment(horizontal="left", vertical="center")
# space-row names are stepped in under their unit header (Excel-native
# indent: formatting only, the cell value stays the exact PDF string)
SPACE_NAME_INDENT = 3
_indented = Alignment(horizontal="left", vertical="center", indent=SPACE_NAME_INDENT)


def _trim_and_load(logo_path: str) -> tuple[BytesIO, int, int]:
    """Crop the logo's surrounding blank margin and return it as PNG bytes.

    Logo files usually carry padding around the artwork (white border, or
    transparent margin). Cropping it first is what makes the mark actually
    read as "filling" the header cell instead of floating in it.
    """
    image = PILImage.open(logo_path)
    image = image.convert("RGBA")

    alpha = image.getchannel("A")
    bbox = alpha.getbbox() if alpha.getextrema()[0] < 255 else None
    if bbox is None:  # opaque image: trim the uniform border colour instead
        rgb = image.convert("RGB")
        background = PILImage.new("RGB", rgb.size, rgb.getpixel((0, 0)))
        bbox = ImageChops.difference(rgb, background).getbbox()
    if bbox:  # None means "entirely uniform" — keep the image as-is
        image = image.crop(bbox)

    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer, image.width, image.height


def _place_logo(ws, logo_path: str) -> None:
    """Embed the user's company logo, trimmed and centered, in merged G1:N1.

    Fit is "contain": the logo is scaled to the largest size that fits the
    cell with its aspect ratio intact — never stretched or cropped, which
    would deform the brand mark. Small logos are enlarged up to
    _LOGO_MAX_UPSCALE so they still read at header size.

    Falls back to the file name as text if the image cannot be loaded, so a
    bad image never costs the engineer the whole export.
    """
    zone_px = sum(_COL_WIDTHS[_LOGO_COL - 1 :]) * _PX_PER_WIDTH
    row_px = _TITLE_ROW_PT * _PT_TO_PX
    max_w = zone_px - 2 * _LOGO_PAD_PX
    max_h = row_px - 2 * _LOGO_PAD_PX

    try:
        buffer, src_w, src_h = _trim_and_load(logo_path)
        image = XLImage(buffer)
        scale = min(max_w / src_w, max_h / src_h, _LOGO_MAX_UPSCALE)
        width, height = max(1, int(src_w * scale)), max(1, int(src_h * scale))
        image.width, image.height = width, height
        image.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=_LOGO_COL - 1,
                colOff=pixels_to_EMU(int((zone_px - width) / 2)),
                row=0,
                rowOff=pixels_to_EMU(int((row_px - height) / 2)),
            ),
            ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
        )
        ws.add_image(image)
    except Exception:
        cell = ws["G1"]
        cell.value = Path(logo_path).stem
        cell.font = Font(bold=True, size=18, color=LOGO_INK)
        cell.alignment = _center


def write_fcu_xlsx(
    details: dict[str, str],
    column_header: list[str],
    data_rows: list[list[str]],
    target: str | Path,
) -> Path:
    target = Path(target)
    wb = Workbook()
    ws = wb.active
    ws.title = "FCU SCHEDULE"

    # ---- rows 1-5: bordered header block ---------------------------------
    for row in range(1, 6):
        for col in range(1, _NUM_COLS + 1):
            ws.cell(row=row, column=col).border = _border

    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = "FCU SCHEDULE"
    title.font = Font(bold=True, size=16)
    title.alignment = _center

    ws.merge_cells("G1:N1")
    ws.row_dimensions[1].height = _TITLE_ROW_PT
    _place_logo(ws, details.get(LOGO_KEY, ""))

    label_font = Font(bold=True, size=10)
    for row, ((l_key, l_label), (r_key, r_label)) in enumerate(
        zip(PROJECT_FIELDS[:4], PROJECT_FIELDS[4:]), start=2
    ):
        left_label = ws.cell(row=row, column=1, value=l_label)
        left_label.font = label_font
        left_label.alignment = _left
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        left_value = ws.cell(row=row, column=2, value=details[l_key].strip())
        left_value.alignment = _left

        right_label = ws.cell(row=row, column=7, value=r_label)
        right_label.font = label_font
        right_label.alignment = _left
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=_NUM_COLS)
        right_value = ws.cell(row=row, column=8, value=details[r_key].strip())
        right_value.alignment = _left

    # ---- row 6: column headers -------------------------------------------
    header_fill = PatternFill("solid", fgColor=NAVY)
    for col, text in enumerate(column_header, start=1):
        cell = ws.cell(row=6, column=col, value=text)
        cell.fill = header_fill
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.alignment = _center
        cell.border = _border
    ws.row_dimensions[6].height = 32

    # ---- data ------------------------------------------------------------
    # unit header rows start flush left; sub-space names are indented
    for r, row_values in enumerate(data_rows, start=7):
        is_unit_row = len(row_values) > 2 and str(row_values[2]).strip()
        for col, value in enumerate(row_values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _border
            cell.alignment = _indented if (col == 1 and not is_unit_row) else _left

    for col, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A7"

    wb.save(target)
    return target
