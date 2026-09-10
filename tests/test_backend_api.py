"""The web API, exercised through FastAPI's TestClient.

These cover the layer the desktop tests cannot reach: the routers, their
error paths, and — the point of the whole port — that a request through HTTP
produces the same bytes the desktop app produces from the same input.

Signing in is mandatory, so the client fixture signs in once; the gate
itself is covered in test_backend_auth.py.
"""

import csv
import io
import json
import tempfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app
from hap_converter.engine import config as hapext_config_mod
from hap_converter.engine import pipeline as hapext_pipeline

REAL_PDF = Path(__file__).parent / "fixtures" / "System Design.pdf"
MAPPING = Path(__file__).resolve().parents[1] / "config" / "mapping.json"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DETAILS = {
    "project": "DCG BREEZE", "project_no": "P-01", "stage": "IFC",
    "discipline": "MEP", "author": "AB", "checked": "CD",
    "revision": "A", "date": "02.09.2026",
}

SCHEDULE_HEADER = [
    "Zone Name / Space Name", "Floor Area (m²)", "Total Coil Load (KW)",
    "Sens Coil Load (KW)", "Air Flow (L/s)",
]
SCHEDULE_ROWS = [
    ["#01-Unit", "172.3", "13.2", "12.6", ""],
    ["#01A-Corridor", "154.4", "", "", "868"],
    ["#01B-Lobby", "23.9", "", "", "218"],
]


@pytest.fixture
def client(engineer_client):
    """Signed in as an ordinary engineer, so these tests exercise the
    routers rather than the gate — and prove a normal seat is enough."""
    return engineer_client


@pytest.fixture
def logo():
    return (Path(__file__).resolve().parents[1] / "company logo.png").read_bytes()


@pytest.fixture
def schedule_csv(tmp_path):
    path = tmp_path / "Tower A.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(SCHEDULE_HEADER)
        writer.writerows(SCHEDULE_ROWS)
    return path


def _convert(client, pdf_path):
    return client.post(
        "/api/hapext/convert",
        files={"pdf": (Path(pdf_path).name, Path(pdf_path).read_bytes(), "application/pdf")},
    ).json()


# ------------------------------------------------------------------- health
def test_health_reports_both_modules_and_the_catalogs(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok"
    assert body["modules"] == ["HAPExt", "AirSizer Pro"]
    assert body["diffusers"] == 5
    assert body["auth"] == "on"   # signing in is always required


# ------------------------------------------------------------------- HAPExt
def test_inspect_returns_the_page_count(client, sample_pdf):
    body = client.post(
        "/api/hapext/inspect",
        files={"pdf": ("sample.pdf", sample_pdf.read_bytes(), "application/pdf")},
    ).json()
    assert body["pages"] == 4          # cover page + three unit pages
    assert body["size"] > 0


def test_inspect_rejects_something_that_is_not_a_pdf(client):
    response = client.post(
        "/api/hapext/inspect", files={"pdf": ("notes.pdf", b"hello", "application/pdf")}
    )
    assert response.status_code == 400
    assert "could not be opened" in response.json()["detail"]


def test_wrong_file_type_is_refused_before_the_engine_sees_it(client):
    response = client.post(
        "/api/hapext/convert", files={"pdf": ("sheet.xlsx", b"PK\x03\x04", XLSX_MIME)}
    )
    assert response.status_code == 400
    assert ".pdf" in response.json()["detail"]


def test_an_empty_upload_is_refused(client):
    response = client.post(
        "/api/hapext/convert", files={"pdf": ("empty.pdf", b"", "application/pdf")}
    )
    assert response.status_code == 400


def test_convert_returns_every_row(client, sample_pdf):
    body = _convert(client, sample_pdf)
    assert body["ok"] is True
    assert body["stats"]["units"] == 3
    assert body["stats"]["spaces"] == 4
    assert len(body["header"]) == 14
    assert len(body["rows"]) == 7                 # 3 unit rows + 4 space rows
    assert body["rows"][0][0] == "#01-9FCorridor1(LIFT)"


def test_validation_failure_blocks_and_names_page_and_field(client, bad_pdf):
    body = _convert(client, bad_pdf)
    assert body["ok"] is False
    assert body["rows"] == [] if "rows" in body else True
    fields = {issue["field"] for issue in body["issues"]}
    assert "Water Flow @9.0 K (L/s)" in fields
    assert "Air Flow (L/s)" in fields
    assert all(issue["page"] == 1 for issue in body["issues"])


@pytest.mark.skipif(not REAL_PDF.is_file(), reason="real report fixture not present")
def test_web_output_is_byte_identical_to_the_desktop_output(client, tmp_path):
    """The whole point of the port: same input, same bytes, either front end."""
    body = _convert(client, REAL_PDF)
    assert body["stats"]["units"] == 212
    assert body["stats"]["spaces"] == 601

    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(body["header"])
    writer.writerows(body["rows"])
    from_web = b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")

    config = hapext_config_mod.load(MAPPING)
    from_desktop = hapext_pipeline.convert(
        str(REAL_PDF), str(tmp_path), config
    ).output_path.read_bytes()

    assert from_web == from_desktop


def test_download_builds_the_xlsx_with_the_logo_embedded(client, sample_pdf, logo):
    body = _convert(client, sample_pdf)
    response = client.post(
        "/api/hapext/download",
        data={
            "payload": json.dumps({"header": body["header"], "rows": body["rows"],
                                   "base_name": "Sample"}),
            "details": json.dumps(DETAILS), "fmt": "xlsx",
        },
        files={"logo": ("logo.png", logo, "image/png")},
    )
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet["A1"].value == "FCU SCHEDULE"
    assert sheet["B2"].value == "DCG BREEZE"
    assert len(sheet._images) == 1                 # the company logo
    assert "Sample.xlsx" in response.headers["content-disposition"]


def test_download_as_csv_carries_the_project_header(client, sample_pdf, logo):
    body = _convert(client, sample_pdf)
    response = client.post(
        "/api/hapext/download",
        data={
            "payload": json.dumps({"header": body["header"], "rows": body["rows"],
                                   "base_name": "Sample"}),
            "details": json.dumps(DETAILS), "fmt": "csv",
        },
        files={"logo": ("logo.png", logo, "image/png")},
    )
    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0][0] == "FCU SCHEDULE"
    assert rows[1][:2] == ["Project:", "DCG BREEZE"]
    assert rows[5] == body["header"]               # 5 header rows, then columns


def test_download_is_refused_when_a_mandatory_detail_is_missing(client, sample_pdf, logo):
    body = _convert(client, sample_pdf)
    incomplete = dict(DETAILS, author="")
    response = client.post(
        "/api/hapext/download",
        data={
            "payload": json.dumps({"header": body["header"], "rows": body["rows"],
                                   "base_name": "Sample"}),
            "details": json.dumps(incomplete), "fmt": "xlsx",
        },
        files={"logo": ("logo.png", logo, "image/png")},
    )
    assert response.status_code == 400
    assert "Author" in response.json()["detail"]


def test_download_with_nothing_to_write_is_refused(client):
    response = client.post(
        "/api/hapext/download",
        data={"payload": json.dumps({"header": [], "rows": []}),
              "details": json.dumps(DETAILS), "fmt": "csv"},
    )
    assert response.status_code == 400


# ----------------------------------------------------------- change request
def _schedule_workbook(header, rows, logo_path, tmp_path):
    from hap_converter.engine.synthesizer import LOGO_KEY
    from hap_converter.engine.xlsx_exporter import write_fcu_xlsx

    target = tmp_path / "previous.xlsx"
    write_fcu_xlsx(dict(DETAILS, **{LOGO_KEY: str(logo_path)}), header, rows, target)
    return target


def test_inspect_schedule_accepts_a_workbook_this_app_produced(client, sample_pdf, tmp_path):
    body = _convert(client, sample_pdf)
    logo_path = Path(__file__).resolve().parents[1] / "company logo.png"
    workbook = _schedule_workbook(body["header"], body["rows"], logo_path, tmp_path)
    response = client.post(
        "/api/hapext/inspect-schedule",
        files={"xlsx": ("previous.xlsx", workbook.read_bytes(), XLSX_MIME)},
    ).json()
    assert response["ok"] is True
    assert response["rows"] == 7
    assert response["columns"] == 14


def test_inspect_schedule_rejects_a_foreign_workbook(client, tmp_path):
    from openpyxl import Workbook

    stray = tmp_path / "stray.xlsx"
    Workbook().save(stray)
    body = client.post(
        "/api/hapext/inspect-schedule",
        files={"xlsx": ("stray.xlsx", stray.read_bytes(), XLSX_MIME)},
    ).json()
    assert body["ok"] is False
    assert "not generated by HAPExt" in body["message"]


def test_change_request_appends_only_the_new_units(client, sample_pdf, tmp_path):
    body = _convert(client, sample_pdf)
    logo_path = Path(__file__).resolve().parents[1] / "company logo.png"
    # a previous schedule holding only the first unit block
    partial = body["rows"][:3]
    workbook = _schedule_workbook(body["header"], partial, logo_path, tmp_path)

    result = client.post(
        "/api/hapext/change-request",
        files={
            "xlsx": ("previous.xlsx", workbook.read_bytes(), XLSX_MIME),
            "pdf": ("sample.pdf", sample_pdf.read_bytes(), "application/pdf"),
        },
    ).json()

    assert result["ok"] is True
    assert result["existing_rows"] == 3
    assert result["stats"]["new_units"] == 2
    assert result["stats"]["total_rows"] == 7
    # the workbook comes back inline so the browser can save it statelessly
    import base64
    saved = base64.b64decode(result["file_b64"])
    assert openpyxl.load_workbook(io.BytesIO(saved)).active["A1"].value == "FCU SCHEDULE"


def test_change_request_with_nothing_new_is_reported_not_appended(client, sample_pdf, tmp_path):
    body = _convert(client, sample_pdf)
    logo_path = Path(__file__).resolve().parents[1] / "company logo.png"
    workbook = _schedule_workbook(body["header"], body["rows"], logo_path, tmp_path)

    result = client.post(
        "/api/hapext/change-request",
        files={
            "xlsx": ("previous.xlsx", workbook.read_bytes(), XLSX_MIME),
            "pdf": ("sample.pdf", sample_pdf.read_bytes(), "application/pdf"),
        },
    ).json()
    assert result["ok"] is False
    assert result["issues"][0]["field"] == "no_changes"


# ---------------------------------------------------------------- AirSizer
def test_config_exposes_the_input_matrix_that_drives_the_form(client):
    body = client.get("/api/airsizer/config").json()
    keys = [d["key"] for d in body["diffusers"]]
    assert keys == ["linear_slot", "linear_bar", "flow_bar", "square", "grille"]

    slot = next(d for d in body["diffusers"] if d["key"] == "linear_slot")
    assert slot["group"] == "A"
    assert slot["inputs"] == ["nc", "slot_width", "no_of_slots"]
    assert slot["options"]["slot_width"] == ["16", "20", "25"]   # per-type override

    assert any(c["locked"] for c in body["result_columns"])      # the name column
    assert "interpolated value" in body["interpolation_remark"]


def test_every_diffuser_serves_its_diagram(client):
    for spec in client.get("/api/airsizer/config").json()["diffusers"]:
        response = client.get(spec["diagram"])
        assert response.status_code == 200, spec["key"]
        assert response.headers["content-type"] == "image/png"


def test_an_unknown_diffuser_diagram_is_a_404(client):
    assert client.get("/api/airsizer/diagram/nonsense").status_code == 404


def test_load_reads_the_schedule_into_subspaces(client, schedule_csv):
    body = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()
    assert [s["name"] for s in body["spaces"]] == [r[0] for r in SCHEDULE_ROWS]
    assert [s["sizable"] for s in body["spaces"]] == [False, True, True]


def test_load_rejects_a_file_that_is_not_a_schedule(client, tmp_path):
    stray = tmp_path / "notes.csv"
    stray.write_text("a,b\n1,2\n", encoding="utf-8")
    response = client.post(
        "/api/airsizer/load", files={"schedule": ("notes.csv", stray.read_bytes(), "text/csv")}
    )
    assert response.status_code == 400
    assert "column header" in response.json()["detail"]


def test_load_rows_carries_a_hapext_run_straight_over(client, sample_pdf):
    converted = _convert(client, sample_pdf)
    body = client.post(
        "/api/airsizer/load-rows",
        json={"header": converted["header"], "rows": converted["rows"],
              "source": "sample.pdf", "base_name": "sample"},
    ).json()
    assert len(body["spaces"]) == len(converted["rows"])
    assert sum(1 for s in body["spaces"] if s["sizable"]) == 4


def test_size_reads_an_exact_catalogue_cell(client, schedule_csv):
    spaces = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()["spaces"]
    target = spaces[2]                                   # 218 L/s

    body = client.post("/api/airsizer/size", json={
        "space": target, "diffuser": "linear_bar",
        "values": {"velocity": "2.5", "nc": "15", "grille_height": "150", "bar_pitch": "6"},
    }).json()

    assert body["ok"] is True
    assert body["interpolated"] is False
    assert body["status"] == "Read"
    assert body["lsm"] == "100"
    assert body["length_m"] == "2.18"
    assert body["pieces"] == 3
    assert body["throw"] == "4.4"


def test_size_flags_an_interpolated_reading(client, schedule_csv):
    spaces = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()["spaces"]
    body = client.post("/api/airsizer/size", json={
        "space": spaces[2], "diffuser": "linear_slot",
        "values": {"nc": "20", "slot_width": "20", "no_of_slots": "3"},
    }).json()
    assert body["ok"] is True
    assert body["interpolated"] is True
    assert body["status"] == "Interpolated"
    assert body["lsm"] == "81.25"


def test_no_valid_selection_is_a_200_with_a_reason(client, schedule_csv):
    """A selection that cannot be made is a normal outcome, not an HTTP error."""
    spaces = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()["spaces"]
    response = client.post("/api/airsizer/size", json={
        "space": spaces[2], "diffuser": "linear_slot",
        "values": {"nc": "15", "slot_width": "16", "no_of_slots": "2"},
    })
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is False
    assert body["status"] == "No valid selection"
    assert "below the quietest" in body["message"]


def test_export_honours_the_column_picker(client, schedule_csv):
    spaces = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()["spaces"]
    inputs = {str(spaces[2]["row"]): {
        "diffuser": "linear_bar",
        "values": {"velocity": "2.5", "nc": "15", "grille_height": "150", "bar_pitch": "6"},
    }}

    response = client.post("/api/airsizer/export", data={"payload": json.dumps({
        "spaces": spaces, "inputs": inputs,
        "visible_columns": ["air_flow", "throw"],
        "base_name": "Tower A", "project_name": "DCG",
    })})
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[4] if c.value] == [
        "Zone Name / Space Name", "Air Flow (L/s)", "Flow Throw (m)",
    ]


def test_export_recomputes_rather_than_trusting_the_browser(client, schedule_csv):
    """The workbook may only ever contain numbers this engine produced."""
    spaces = client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    ).json()["spaces"]
    # no results are posted at all — only the inputs
    inputs = {str(spaces[2]["row"]): {
        "diffuser": "linear_bar",
        "values": {"velocity": "2.5", "nc": "15", "grille_height": "150", "bar_pitch": "6"},
    }}
    response = client.post("/api/airsizer/export", data={"payload": json.dumps({
        "spaces": spaces, "inputs": inputs, "base_name": "Tower A",
    })})
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    values = [c.value for c in sheet[7]]       # header row 4, data 5..7
    assert "100" in values                     # L/s/m the server derived
    assert "2.18 / 3" in values                # length / pieces


def test_export_with_no_rows_is_refused(client):
    response = client.post(
        "/api/airsizer/export", data={"payload": json.dumps({"spaces": []})}
    )
    assert response.status_code == 400


# ------------------------------------------------------------ statelessness
def test_requests_leave_no_temp_files_behind(client, sample_pdf, schedule_csv):
    """Render's disk is ephemeral; nothing may accumulate between requests."""
    scratch = Path(tempfile.gettempdir())
    before = {p.name for p in scratch.glob("maec_*")}

    _convert(client, sample_pdf)
    client.post(
        "/api/airsizer/load",
        files={"schedule": ("Tower A.csv", schedule_csv.read_bytes(), "text/csv")},
    )
    client.post(
        "/api/hapext/inspect",
        files={"pdf": ("sample.pdf", sample_pdf.read_bytes(), "application/pdf")},
    )

    assert {p.name for p in scratch.glob("maec_*")} == before
