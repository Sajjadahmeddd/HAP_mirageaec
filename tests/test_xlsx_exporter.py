import openpyxl
import pytest
from PIL import Image

from hap_converter.engine.xlsx_exporter import write_fcu_xlsx

EMU_PER_PX = 9525

BASE_DETAILS = {
    "project": "Avarra by Palace",
    "project_no": "MLD",
    "stage": "100 % DD",
    "discipline": "MEP",
    "author": "SA",
    "checked": "FA",
    "revision": "2",
    "date": "20.04.2026",
}


@pytest.fixture
def logo(tmp_path):
    path = tmp_path / "company logo.png"
    Image.new("RGBA", (600, 200), (20, 60, 90, 255)).save(path)
    return path


@pytest.fixture
def DETAILS(logo):
    return dict(BASE_DETAILS, logo_path=str(logo))

HEADER = [f"Col {i}" for i in range(1, 15)]
DATA = [
    ["#01-Corridor", "178.2", "9.2", "8.6", "", "23.5 / 16.7", "13.2 / 12.5",
     "0.24", "51.63", "", "9.2", "", "", ""],
    ["#01-Corridor", "178.2", "", "", "689", "", "", "", "", "", "", "", "", ""],
]


def _load(tmp_path, details):
    target = tmp_path / "schedule.xlsx"
    write_fcu_xlsx(details, HEADER, DATA, target)
    return openpyxl.load_workbook(target)["FCU SCHEDULE"]


def test_row1_is_two_merged_zones(tmp_path, DETAILS):
    ws = _load(tmp_path, DETAILS)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:F1" in merged  # FCU SCHEDULE, no column lines inside
    assert "G1:N1" in merged  # logo zone, no column lines inside
    assert ws["A1"].value == "FCU SCHEDULE"


def test_company_logo_image_embedded_in_logo_zone(tmp_path, DETAILS):
    ws = _load(tmp_path, DETAILS)
    assert len(ws._images) == 1  # the user's logo image, not text
    image = ws._images[0]
    assert ws["G1"].value is None  # no placeholder text left behind
    assert image.anchor._from.col == 6  # column G (0-based)
    assert image.anchor._from.row == 0  # row 1
    # displayed size (anchor extent, in EMU) is scaled to fit the merged
    # zone with the 600x200 source aspect preserved
    px_w = image.anchor.ext.cx / EMU_PER_PX
    px_h = image.anchor.ext.cy / EMU_PER_PX
    assert px_h <= 88 and px_w <= 613
    assert abs(px_w / px_h - 3.0) < 0.05
    # centered inside the zone rather than jammed against the left edge
    assert image.anchor._from.colOff > 0


def test_unreadable_logo_falls_back_to_text(tmp_path):
    broken = tmp_path / "Acme Corp.png"
    broken.write_text("not really an image")
    ws = _load(tmp_path, dict(BASE_DETAILS, logo_path=str(broken)))
    assert ws._images == []
    assert ws["G1"].value == "Acme Corp"  # export still succeeds


def test_detail_rows_merged_values(tmp_path, DETAILS):
    ws = _load(tmp_path, DETAILS)
    merged = {str(r) for r in ws.merged_cells.ranges}
    for row in range(2, 6):
        assert f"B{row}:F{row}" in merged
        assert f"H{row}:N{row}" in merged
    assert ws["A2"].value == "Project:" and ws["B2"].value == "Avarra by Palace"
    assert ws["G4"].value == "Revision:" and ws["H4"].value == "2"
    assert ws["G5"].value == "Date:" and ws["H5"].value == "20.04.2026"


def test_header_and_data_rows(tmp_path, DETAILS):
    ws = _load(tmp_path, DETAILS)
    assert [ws.cell(row=6, column=c).value for c in range(1, 15)] == HEADER
    assert ws["A7"].value == "#01-Corridor"
    assert ws["B7"].value == "178.2"  # string-exact, zero rounding
    assert ws["E8"].value == "689"
    assert ws.cell(row=7, column=5).value in ("", None)  # blanks stay blank
