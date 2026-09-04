"""Export and project persistence.

The export must contain exactly the columns the engineer left visible on the
Review screen, and must never overwrite an earlier file.
"""

import openpyxl
import pytest

from hap_converter.airsizer.engine import export, project as project_mod
from hap_converter.airsizer.engine.models import SizingInput, Space
from hap_converter.airsizer.engine.pipeline import size_space
from hap_converter.airsizer.engine.project import Project


@pytest.fixture
def sized(air_config):
    """Two subspaces under one unit header: one sized, one left unsized."""
    spaces = [
        Space(name="#01-9F-Corridor1", floor_area="172.3", total_coil="13.2",
              sens_coil="12.6", air_flow="", is_unit=True, row=7),
        Space(name="#01A-9FCorridor1", floor_area="154.4", air_flow="868", row=8),
        Space(name="#01B-9FElec. Room", floor_area="6.1", air_flow="115", row=9),
    ]
    inputs = {
        8: SizingInput("linear_bar", {"velocity": "2.5", "nc": "15",
                                      "grille_height": "150", "bar_pitch": "6"})
    }
    results = {row: size_space(next(s for s in spaces if s.row == row), value, air_config)
               for row, value in inputs.items()}
    return spaces, inputs, results


def test_default_columns_come_from_the_config(air_config):
    columns = export.visible_columns(air_config, None)
    assert [c.key for c in columns] == [
        "name", "floor_area", "total_coil", "air_flow", "lsm", "length", "remarks",
    ]


def test_column_picker_choice_is_honoured_and_name_stays_locked(air_config):
    columns = export.visible_columns(air_config, ["air_flow", "throw"])
    keys = [c.key for c in columns]
    assert keys == ["name", "air_flow", "throw"]      # name is locked on
    assert "floor_area" not in keys


def test_rows_follow_the_schedule_and_blank_the_unsized(sized, air_config):
    spaces, inputs, results = sized
    columns = export.visible_columns(air_config, None)
    header, rows = export.build_rows(spaces, inputs, results, air_config, columns)

    assert header == ["Zone Name / Space Name", "Floor Area (m²)",
                      "Total Coil Load (KW)", "Air Flow (L/s)",
                      "Air Flow (L/s/m)", "Air Outlet Length (m) / Nos",
                      "Remarks"]
    assert len(rows) == 3
    assert rows[0][:4] == ["#01-9F-Corridor1", "172.3", "13.2", ""]
    assert rows[0][4:] == ["", "", ""]                # unit rows carry no sizing
    assert rows[1] == ["#01A-9FCorridor1", "154.4", "", "868", "100", "8.68 / 9", ""]
    assert rows[2][4:] == ["", "", ""]                # not sized yet


def test_remarks_carry_the_interpolation_sentence(sized, air_config):
    """An interpolated row gets the exact sentence the sizing panel shows."""
    import dataclasses

    spaces, inputs, results = sized
    results = {8: dataclasses.replace(results[8], interpolated=True)}
    columns = export.visible_columns(air_config, ["remarks"])
    header, rows = export.build_rows(spaces, inputs, results, air_config, columns)

    assert header == ["Zone Name / Space Name", "Remarks"]
    assert air_config.interpolation_remark            # the config carries it
    assert rows[1][1] == air_config.interpolation_remark
    assert rows[0][1] == ""                           # unit header row
    assert rows[2][1] == ""                           # unsized row


def test_remarks_blank_when_the_reading_was_exact(sized, air_config):
    spaces, inputs, results = sized
    assert results[8].interpolated is False           # fixture is an exact hit
    columns = export.visible_columns(air_config, ["remarks"])
    _, rows = export.build_rows(spaces, inputs, results, air_config, columns)
    assert rows[1][1] == ""


def test_row_state_drives_the_highlighting(sized, air_config):
    spaces, _, results = sized
    assert export.row_state(spaces[0], results.get(7)) == "unit"
    assert export.row_state(spaces[1], results.get(8)) == "ok"
    assert export.row_state(spaces[2], results.get(9)) == "unsized"


def test_export_writes_only_the_visible_columns(sized, air_config, tmp_path):
    spaces, inputs, results = sized
    columns = export.visible_columns(air_config, ["air_flow", "throw"])
    path = export.write_xlsx(
        spaces, inputs, results, air_config, columns, tmp_path, "sized", "Sample Tower"
    )
    sheet = openpyxl.load_workbook(path).active
    assert [c.value for c in sheet[4][:3]] == [
        "Zone Name / Space Name", "Air Flow (L/s)", "Flow Throw (m)"
    ]
    # 868 L/s at 100 L/s/m is an 8.68 m run: the 6-10 m band throws x1.15
    assert sheet.cell(row=6, column=3).value == "4.6"
    assert sheet.cell(row=4, column=4).value is None      # nothing beyond the three


def test_the_project_block_heads_the_sheet_when_details_are_given(sized, air_config, tmp_path):
    """The same nine inputs the HAPExt schedule carries, laid out the same way,
    so the two documents read as one set."""
    from PIL import Image

    logo = tmp_path / "client.png"
    Image.new("RGBA", (400, 140), (10, 80, 50, 255)).save(logo)
    details = {
        "project": "Avarra by Palace", "project_no": "MLD", "stage": "100 % DD",
        "discipline": "MEP", "author": "SA", "checked": "FA", "revision": "2",
        "date": "20.04.2026", "logo_path": str(logo),
    }
    spaces, inputs, results = sized
    columns = export.visible_columns(air_config, None)
    path = export.write_xlsx(spaces, inputs, results, air_config, columns,
                             tmp_path, "sized", "Tower", details)

    sheet = openpyxl.load_workbook(path).active
    assert sheet["A1"].value == "AIR DIFFUSER SIZING"
    assert (sheet["A2"].value, sheet["B2"].value) == ("Project:", "Avarra by Palace")
    assert (sheet["G5"].value, sheet["H5"].value) == ("Date:", "20.04.2026")
    assert len(sheet._images) == 1                       # the client's logo
    assert sheet.cell(row=6, column=1).value == "Zone Name / Space Name"
    assert sheet.cell(row=7, column=1).value == "#01-9F-Corridor1"


def test_without_details_the_sheet_keeps_its_plain_heading(sized, air_config, tmp_path):
    spaces, inputs, results = sized
    columns = export.visible_columns(air_config, None)
    path = export.write_xlsx(spaces, inputs, results, air_config, columns,
                             tmp_path, "plain", "Tower")
    sheet = openpyxl.load_workbook(path).active
    assert sheet["A1"].value == "AIR DIFFUSER SIZING"
    assert sheet["A2"].value == "Tower"                  # the old subtitle
    assert sheet.cell(row=4, column=1).value == "Zone Name / Space Name"
    assert sheet._images == []


def test_export_auto_versions_instead_of_overwriting(sized, air_config, tmp_path):
    spaces, inputs, results = sized
    columns = export.visible_columns(air_config, None)
    first = export.write_xlsx(spaces, inputs, results, air_config, columns, tmp_path, "sized")
    second = export.write_xlsx(spaces, inputs, results, air_config, columns, tmp_path, "sized")
    third = export.write_xlsx(spaces, inputs, results, air_config, columns, tmp_path, "sized")
    assert first.name == "sized.xlsx"
    assert second.name == "sized_v1.xlsx"
    assert third.name == "sized_v2.xlsx"
    assert first.exists() and second.exists()


# ---------------------------------------------------------------- projects
def test_project_round_trips_through_json(tmp_path):
    original = Project(name="Office Tower", source_path="C:/schedule.xlsx",
                       visible_columns=["name", "air_flow"])
    original.set_sizing(8, SizingInput("linear_bar", {"nc": "15", "bar_pitch": "6"}))
    path = project_mod.save(original, tmp_path / "office.airsizer.json")

    reopened = project_mod.load(path)
    assert reopened.name == "Office Tower"
    assert reopened.source_path == "C:/schedule.xlsx"
    assert reopened.visible_columns == ["name", "air_flow"]
    assert reopened.sized_count == 1
    restored = reopened.sizing_for(8)
    assert restored.diffuser == "linear_bar"
    assert restored.values == {"nc": "15", "bar_pitch": "6"}
    assert reopened.sizing_for(99) is None


def test_saving_stamps_created_and_updated(tmp_path):
    saved = Project(name="Stamped")
    project_mod.save(saved, tmp_path / "stamped.json")
    assert saved.created and saved.updated


def test_clearing_a_sizing_removes_it(tmp_path):
    saved = Project(name="Cleared")
    saved.set_sizing(8, SizingInput("square", {"nc": "25"}))
    saved.clear_sizing(8)
    assert saved.sized_count == 0


def test_recent_projects_puts_the_newest_first(tmp_path):
    index = tmp_path / "recent.json"
    for name in ("first", "second"):
        saved = Project(name=name, path=tmp_path / f"{name}.json")
        saved.updated = "2026-08-24T10:00:00"
        project_mod.RecentProjects(index).add(saved)

    recents = project_mod.RecentProjects(index)
    assert [e.name for e in recents.top()] == ["second", "first"]
    recents.remove(str(tmp_path / "second.json"))
    assert [e.name for e in project_mod.RecentProjects(index).top()] == ["first"]


def test_a_corrupt_recents_index_starts_empty(tmp_path):
    index = tmp_path / "recent.json"
    index.write_text("{not json", encoding="utf-8")
    assert project_mod.RecentProjects(index).entries == []
