"""Row assembly + the sized-schedule Excel writer.

`build_rows` is the single place that turns (schedule rows, chosen inputs,
sizing results) into a table. The Review Results screen renders exactly what
this returns, and the exporter writes exactly what it returns — so what the
engineer approves on screen is what lands in the file, including which
columns are visible.

Output filenames auto-version through module 1's `exporter.versioned_path`
(name.xlsx, name_v1.xlsx, ...) so a re-export never overwrites.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from hap_converter.engine.exporter import versioned_path

from .config import Config, ResultColumn
from .models import Space, SizingInput, SizingResult

NAVY = "141B4D"
GRID = "9A9A9A"
INTERPOLATED_FILL = "FFF3D6"   # amber tint: value read between catalogue entries
FAILED_FILL = "FDEDED"         # red tint: no valid selection
SPACE_NAME_INDENT = 3

_thin = Side(style="thin", color=GRID)
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
_left = Alignment(horizontal="left", vertical="center")
_indented = Alignment(horizontal="left", vertical="center", indent=SPACE_NAME_INDENT)


def visible_columns(config: Config, keys: list[str] | None) -> list[ResultColumn]:
    """The result columns to show, in config order. None = the defaults."""
    if keys is None:
        return [c for c in config.result_columns if c.default]
    chosen = set(keys)
    return [c for c in config.result_columns if c.key in chosen or c.locked]


def cell_value(
    column_key: str,
    space: Space,
    sizing_input: SizingInput | None,
    result: SizingResult | None,
    config: Config,
) -> str:
    """One cell of the review/export table."""
    if column_key == "name":
        return space.name
    if column_key == "floor_area":
        return space.floor_area
    if column_key == "total_coil":
        return space.total_coil
    if column_key == "sens_coil":
        return space.sens_coil
    if column_key == "air_flow":
        return space.air_flow

    # sizing columns are blank on unit header rows and on unsized subspaces
    if space.is_unit or result is None:
        return ""

    if column_key == "diffuser":
        if not sizing_input:
            return ""
        try:
            return config.diffuser(sizing_input.diffuser).label
        except Exception:
            return sizing_input.diffuser
    if column_key == "status":
        return result.status
    if not result.ok:
        return ""

    if column_key == "lsm":
        return result.lsm
    if column_key == "length":
        if result.group == "A":
            return f"{result.length_m} / {result.pieces}" if result.length_m else ""
        return str(result.outlets) if result.outlets else ""
    if column_key == "throw":
        return result.throw
    if column_key == "size":
        return result.size
    if column_key == "nc_out":
        return result.nc
    if column_key == "velocity_out":
        return result.velocity
    if column_key == "pt":
        return result.pt
    if column_key == "interpolated":
        return "Yes" if result.interpolated else ""
    if column_key == "remarks":
        # the same sentence the sizing panel shows for a value read between
        # two catalogue entries; one source of truth in input_matrix.json
        return config.interpolation_remark if result.interpolated else ""
    return ""


def build_rows(
    spaces: list[Space],
    inputs: dict[int, SizingInput],
    results: dict[int, SizingResult],
    config: Config,
    columns: list[ResultColumn],
) -> tuple[list[str], list[list[str]]]:
    """Header labels + one row per schedule row, in schedule order."""
    header = [column.label for column in columns]
    rows = [
        [
            cell_value(column.key, space, inputs.get(space.row), results.get(space.row), config)
            for column in columns
        ]
        for space in spaces
    ]
    return header, rows


def row_state(space: Space, result: SizingResult | None) -> str:
    """"unit" | "ok" | "interpolated" | "failed" | "unsized" — drives tinting."""
    if space.is_unit:
        return "unit"
    if result is None:
        return "unsized"
    if not result.ok:
        return "failed"
    return "interpolated" if result.interpolated else "ok"


def write_xlsx(
    spaces: list[Space],
    inputs: dict[int, SizingInput],
    results: dict[int, SizingResult],
    config: Config,
    columns: list[ResultColumn],
    output_dir: str | Path,
    base_name: str,
    project_name: str = "",
) -> Path:
    """Write the sized schedule, auto-versioned, and return the path."""
    header, rows = build_rows(spaces, inputs, results, config, columns)
    target = versioned_path(output_dir, base_name, suffix=".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AIR DIFFUSER SIZING"

    title = sheet.cell(row=1, column=1, value="AIR DIFFUSER SIZING")
    title.font = Font(bold=True, size=14)
    title.alignment = _left
    if project_name:
        subtitle = sheet.cell(row=2, column=1, value=project_name)
        subtitle.font = Font(size=10, color="666666")

    header_row = 4
    header_fill = PatternFill("solid", fgColor=NAVY)
    for index, label in enumerate(header, start=1):
        cell = sheet.cell(row=header_row, column=index, value=label)
        cell.fill = header_fill
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.alignment = _center
        cell.border = _border
    sheet.row_dimensions[header_row].height = 32

    interpolated = PatternFill("solid", fgColor=INTERPOLATED_FILL)
    failed = PatternFill("solid", fgColor=FAILED_FILL)
    for offset, (space, values) in enumerate(zip(spaces, rows), start=header_row + 1):
        state = row_state(space, results.get(space.row))
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=index, value=value)
            cell.border = _border
            cell.alignment = _indented if (index == 1 and not space.is_unit) else _left
            if state == "interpolated":
                cell.fill = interpolated
            elif state == "failed":
                cell.fill = failed
            if space.is_unit:
                cell.font = Font(bold=True, size=10)

    widths = {"name": 34, "length": 20, "size": 16, "diffuser": 22, "status": 18}
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column.key, 14)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    workbook.save(target)
    return target
